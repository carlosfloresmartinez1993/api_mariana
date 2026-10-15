"""Carga masiva de entradas y salidas de un proyecto desde el Excel del cliente.

Formato (fijo): una hoja con las columnas FECHA, CONCEPTO, CARGO y ABONO.
  ABONO = entrada, CARGO = salida. Cada fila lleva uno de los dos.
Lo que el Excel no trae (empresa, cuentas, categorías, % de gasto) se elige al cargar.

Todo o nada: si el archivo o alguna fila tiene un error no se registra nada. Primero se registran todas las
entradas y después todas las salidas, en una sola transacción.
"""

import io
import unicodedata
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.errors import ReglaNegocio
from app.core.tiempo import hoy
from app.models import AccionBitacora, Categoria, Cuenta, Entrada, Forma, Proyecto, Salida, Usuario
from app.services.acceso import obtener
from app.services.bitacora import instantanea, registrar
from app.services.movimientos import CAMPOS_ENTRADA, CAMPOS_SALIDA, vincular
from app.services.saldos import bloquear, saldos_cuentas, validar_no_negativos

COLUMNAS = ("FECHA", "CONCEPTO", "CARGO", "ABONO")
MAX_FILAS = 5000
MAXIMO = Decimal("999999999999.99")


@dataclass
class Fila:
    numero: int  # número de fila en el Excel
    fecha: date
    concepto: str | None
    tipo: str  # "entrada" | "salida"
    monto: Decimal


@dataclass
class Error:
    fila: int | None
    columna: str | None
    mensaje: str


@dataclass
class Opciones:
    empresa_id: uuid.UUID
    cuenta_entradas_id: uuid.UUID
    cuenta_salidas_id: uuid.UUID
    porcentaje_gasto: Decimal = Decimal("0")
    categoria_entradas_id: uuid.UUID | None = None
    categoria_salidas_id: uuid.UUID | None = None
    requiere_factura_entradas: bool = False
    requiere_factura_salidas: bool = False


@dataclass
class Resultado:
    errores: list[Error] = field(default_factory=list)
    filas: list[Fila] = field(default_factory=list)


def _normalizar(texto: Any) -> str:
    if texto is None:
        return ""
    sin_acentos = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return " ".join(sin_acentos.upper().split())


def _monto(valor: Any) -> Decimal | None:
    """None si la celda está vacía. Acepta números y textos como "$1,234.50". Lanza ValueError si no es válido."""
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return None
    if isinstance(valor, bool):
        raise ValueError
    if isinstance(valor, (int, float, Decimal)):
        numero = Decimal(str(valor))
    else:
        texto = str(valor).strip().replace("$", "").replace(",", "").replace(" ", "")
        try:
            numero = Decimal(texto)
        except InvalidOperation:
            raise ValueError from None
    return numero.quantize(Decimal("0.01"))


def _fecha(valor: Any) -> date:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    raise ValueError


def leer(contenido: bytes) -> Resultado:
    """Revisa la estructura y cada fila del Excel. No toca la base de datos."""
    resultado = Resultado()
    if len(contenido) > settings.MAX_DOCUMENTO_MB * 1024 * 1024 * 2:
        resultado.errores.append(Error(None, None, "El archivo es demasiado grande"))
        return resultado
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - cualquier archivo que no sea un .xlsx válido
        resultado.errores.append(Error(None, None, "No es un archivo de Excel válido (.xlsx)"))
        return resultado
    hoja = libro.worksheets[0]
    filas = list(hoja.iter_rows(values_only=True))
    libro.close()

    # Encabezado: la primera fila (de las 10 primeras) que tenga las cuatro columnas.
    encabezado, posiciones = None, {}
    for i, fila in enumerate(filas[:10]):
        nombres = [_normalizar(v) for v in fila]
        if all(c in nombres for c in COLUMNAS):
            encabezado, posiciones = i, {c: nombres.index(c) for c in COLUMNAS}
            break
    if encabezado is None:
        resultado.errores.append(Error(None, None, "El archivo no tiene la estructura esperada: la primera hoja debe "
                                                   "tener las columnas FECHA, CONCEPTO, CARGO y ABONO"))
        return resultado

    hoy_ = hoy()
    for i, fila in enumerate(filas[encabezado + 1:], start=encabezado + 2):
        celdas = {c: (fila[p] if p < len(fila) else None) for c, p in posiciones.items()}
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in celdas.values()):
            continue  # fila vacía
        if len(resultado.filas) >= MAX_FILAS:
            resultado.errores.append(Error(i, None, f"Máximo {MAX_FILAS} movimientos por archivo"))
            break
        errores_fila: list[Error] = []
        try:
            fecha = _fecha(celdas["FECHA"])
            if fecha > hoy_:
                errores_fila.append(Error(i, "FECHA", f"La fecha {fecha:%d/%m/%Y} es futura"))
        except ValueError:
            fecha = None
            errores_fila.append(Error(i, "FECHA", "Fecha vacía o no válida (use dd/mm/aaaa)"))
        montos: dict[str, Decimal | None] = {}
        for columna in ("CARGO", "ABONO"):
            try:
                montos[columna] = _monto(celdas[columna])
            except ValueError:
                montos[columna] = None
                errores_fila.append(Error(i, columna, f"«{celdas[columna]}» no es un monto válido"))
        cargo, abono = montos["CARGO"], montos["ABONO"]
        if cargo is not None and abono is not None and cargo and abono:
            errores_fila.append(Error(i, "CARGO / ABONO", "La fila tiene cargo y abono; debe tener solo uno"))
        elif not cargo and not abono and not any(e.columna in ("CARGO", "ABONO") for e in errores_fila):
            errores_fila.append(Error(i, "CARGO / ABONO", "La fila no tiene cargo ni abono"))
        for columna, monto in montos.items():
            if monto is not None and monto < 0:
                errores_fila.append(Error(i, columna, "El monto no puede ser negativo"))
            if monto is not None and monto > MAXIMO:
                errores_fila.append(Error(i, columna, "El monto es demasiado grande"))
        concepto = " ".join(str(celdas["CONCEPTO"]).split())[:255] if celdas["CONCEPTO"] is not None else None
        concepto = concepto or None
        tipo = "salida" if cargo else "entrada"
        if tipo == "salida" and not concepto:
            errores_fila.append(Error(i, "CONCEPTO", "Los cargos (salidas) necesitan concepto"))
        if errores_fila:
            resultado.errores += errores_fila
        else:
            resultado.filas.append(Fila(i, fecha, concepto, tipo, cargo if tipo == "salida" else abono))
    if not resultado.filas and not resultado.errores:
        resultado.errores.append(Error(None, None, "El archivo no tiene movimientos"))
    return resultado


def _cuenta(db: Session, cuenta_id: uuid.UUID, empresa_id: uuid.UUID, para: str) -> Cuenta:
    cuenta = obtener(db, Cuenta, cuenta_id, f"Cuenta para {para}")
    if cuenta.empresa_id != empresa_id:
        raise ReglaNegocio(f"La cuenta para {para} no es de la empresa elegida")
    if not cuenta.activa or not cuenta.empresa.activa:
        raise ReglaNegocio(f"La cuenta para {para} («{cuenta.nombre}») o su empresa está desactivada")
    return cuenta


def _categoria(db: Session, categoria_id: uuid.UUID | None) -> Categoria | None:
    if categoria_id is None:
        return None
    categoria = obtener(db, Categoria, categoria_id, "Categoría")
    if not categoria.activa:
        raise ReglaNegocio(f"La categoría «{categoria.nombre}» está desactivada")
    return categoria


def _forma(cuenta: Cuenta) -> Forma:
    return Forma.EFECTIVO if cuenta.tipo == "caja" else Forma.TRANSFERENCIA


def _neto(monto: Decimal, porcentaje: Decimal) -> tuple[Decimal, Decimal]:
    gasto = (monto * porcentaje / 100).quantize(Decimal("0.01"))
    return gasto, monto - gasto


def vista_previa(db: Session, proyecto: Proyecto, contenido: bytes, op: Opciones) -> dict[str, Any]:
    """Revisa el archivo y calcula cómo quedarían las cuentas, sin registrar nada."""
    if not proyecto.activo:
        raise ReglaNegocio(f"El proyecto «{proyecto.nombre}» está desactivado")
    cuenta_e = _cuenta(db, op.cuenta_entradas_id, op.empresa_id, "las entradas (abonos)")
    cuenta_s = _cuenta(db, op.cuenta_salidas_id, op.empresa_id, "las salidas (cargos)")
    cat_e, cat_s = _categoria(db, op.categoria_entradas_id), _categoria(db, op.categoria_salidas_id)
    resultado = leer(contenido)
    entradas = [f for f in resultado.filas if f.tipo == "entrada"]
    salidas = [f for f in resultado.filas if f.tipo == "salida"]

    bruto = sum((f.monto for f in entradas), Decimal("0.00"))
    gastos = sum((_neto(f.monto, op.porcentaje_gasto)[0] for f in entradas), Decimal("0.00"))
    total_salidas = sum((f.monto for f in salidas), Decimal("0.00"))
    antes = saldos_cuentas(db, list({cuenta_e.id, cuenta_s.id}))
    despues = dict(antes)
    despues[cuenta_e.id] += bruto - gastos
    despues[cuenta_s.id] -= total_salidas
    errores = list(resultado.errores)
    for cuenta in {cuenta_e, cuenta_s}:
        if not errores and despues[cuenta.id] < 0:
            errores.append(Error(None, None, f"La cuenta «{cuenta.nombre}» quedaría en {despues[cuenta.id]:,.2f} "
                                             f"{cuenta.moneda}: no alcanza para pagar las salidas"))
    return {
        "valido": not errores,
        "errores": [e.__dict__ for e in errores],
        "proyecto": proyecto.nombre,
        "entradas": len(entradas),
        "salidas": len(salidas),
        "entradas_bruto": bruto,
        "entradas_gasto": gastos,
        "entradas_neto": bruto - gastos,
        "total_salidas": total_salidas,
        "cuenta_entradas": {"nombre": cuenta_e.nombre, "moneda": cuenta_e.moneda, "forma": _forma(cuenta_e)},
        "cuenta_salidas": {"nombre": cuenta_s.nombre, "moneda": cuenta_s.moneda, "forma": _forma(cuenta_s)},
        "categoria_entradas": cat_e.nombre if cat_e else None,
        "categoria_salidas": cat_s.nombre if cat_s else None,
        "saldos": [{"cuenta": c.nombre, "moneda": c.moneda, "antes": antes[c.id], "despues": despues[c.id]}
                   for c in sorted({cuenta_e, cuenta_s}, key=lambda c: c.nombre)],
        "filas": [{"fila": f.numero, "fecha": f.fecha, "concepto": f.concepto, "tipo": f.tipo, "monto": f.monto}
                  for f in entradas + salidas],
    }


def cargar(db: Session, proyecto: Proyecto, contenido: bytes, nombre_archivo: str | None, op: Opciones,
           usuario: Usuario) -> dict[str, Any]:
    """Registra todo en una sola transacción: primero las entradas, después las salidas."""
    previa = vista_previa(db, proyecto, contenido, op)
    if not previa["valido"]:
        raise ReglaNegocio("El archivo tiene errores; revise la vista previa. No se registró nada.")
    cuenta_e, cuenta_s = db.get(Cuenta, op.cuenta_entradas_id), db.get(Cuenta, op.cuenta_salidas_id)
    bloquear(db, proyecto_ids={proyecto.id}, cuenta_ids={cuenta_e.id, cuenta_s.id})
    vincular(db, cuenta_e.empresa, proyecto, usuario, automatico=True)
    origen = {"origen": "carga masiva", "archivo": nombre_archivo}
    filas = leer(contenido).filas
    creados: list[tuple[str, Any]] = []
    for f in (f for f in filas if f.tipo == "entrada"):
        creados.append(("Entrada", Entrada(
            proyecto_id=proyecto.id, empresa_id=op.empresa_id, cuenta_id=cuenta_e.id, moneda=cuenta_e.moneda,
            fecha=f.fecha, forma=_forma(cuenta_e), monto=f.monto, porcentaje_gasto=op.porcentaje_gasto,
            concepto=f.concepto, categoria_id=op.categoria_entradas_id,
            requiere_factura=op.requiere_factura_entradas, creado_por=usuario.id)))
    for f in (f for f in filas if f.tipo == "salida"):
        creados.append(("Salida", Salida(
            proyecto_id=proyecto.id, empresa_id=op.empresa_id, cuenta_id=cuenta_s.id, moneda=cuenta_s.moneda,
            fecha=f.fecha, forma=_forma(cuenta_s), monto=f.monto, concepto=f.concepto,
            categoria_id=op.categoria_salidas_id, requiere_factura=op.requiere_factura_salidas,
            creado_por=usuario.id)))
    for _, mov in creados:
        db.add(mov)
    db.flush()
    validar_no_negativos(db, cuenta_ids={cuenta_e.id, cuenta_s.id})
    for entidad, mov in creados:
        db.refresh(mov)  # monto_gasto y monto_neto calculados por la base de datos
        campos = CAMPOS_ENTRADA if entidad == "Entrada" else CAMPOS_SALIDA
        registrar(db, usuario_id=usuario.id, entidad=entidad, entidad_id=mov.id, accion=AccionBitacora.CREAR,
                  detalle={"despues": instantanea(mov, campos), **origen})
    db.commit()
    return {"entradas": previa["entradas"], "salidas": previa["salidas"]}


def plantilla() -> bytes:
    """Excel vacío con el formato esperado y una hoja de instrucciones."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Movimientos"
    hoja.append(list(COLUMNAS))
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="000000")
    for letra, ancho in zip("ABCD", (14, 50, 16, 16)):
        hoja.column_dimensions[letra].width = ancho
    ayuda = libro.create_sheet("Instrucciones")
    for linea in (
        "Carga masiva de entradas y salidas",
        "",
        "Llene la hoja «Movimientos» (no cambie los nombres de las columnas):",
        "  FECHA: dd/mm/aaaa, no futura.",
        "  CONCEPTO: descripción del movimiento (obligatorio en los cargos).",
        "  ABONO: dinero que ENTRA al proyecto (entrada).",
        "  CARGO: dinero que SALE del proyecto (salida).",
        "  Cada fila lleva un cargo o un abono, no los dos.",
        "",
        "Antes de cargar, en el sistema deben existir:",
        "  - La empresa a la que va el proyecto.",
        "  - Las cuentas de esa empresa donde entra y de donde sale el dinero (con su moneda).",
        "  - Opcional: las categorías que quiera asignar.",
        "",
        "Al cargar se eligen la empresa, las cuentas, el % de gasto y las categorías; luego se revisa una vista previa.",
        "Si alguna fila tiene un error no se registra nada. Se registran primero las entradas y después las salidas.",
    ):
        ayuda.append([linea])
    ayuda["A1"].font = Font(bold=True, size=14)
    ayuda.column_dimensions["A"].width = 110
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
