import csv
import io
from html import escape
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from enum import StrEnum
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.tiempo import ahora


class Formato(StrEnum):
    JSON = "json"
    CSV = "csv"
    XLSX = "xlsx"
    PDF = "pdf"


@dataclass
class Reporte:
    titulo: str
    columnas: list[tuple[str, str]]  # (clave, etiqueta)
    filas: list[dict[str, Any]]
    # Una fila de totales por moneda (MXN y USD nunca se suman entre sí).
    totales: list[dict[str, Any]] = field(default_factory=list)
    parametros: dict[str, Any] = field(default_factory=dict)
    # Nombre de quien genera el reporte; se imprime al final.
    elaborado_por: str | None = None
    # "tabla" (columnas y filas) o "resumen" (estado de ingresos y egresos: cada fila trae su `estilo`).
    estilo: str = "tabla"


MEDIA = {
    Formato.CSV: "text/csv; charset=utf-8",
    Formato.XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    Formato.PDF: "application/pdf",
}


def _texto(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, Decimal):
        return f"{valor:,.2f}"
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M")
    if isinstance(valor, date):
        return valor.isoformat()
    return str(valor)


def _filas_totales(reporte: Reporte) -> list[list[Any]]:
    filas = []
    for totales in reporte.totales:
        fila = [totales.get(clave) for clave, _ in reporte.columnas]
        if fila and fila[0] is None:
            fila[0] = "TOTAL"
        filas.append(fila)
    return filas


def _pie(reporte: Reporte) -> str | None:
    if not reporte.elaborado_por:
        return None
    return f"Elaborado por: {reporte.elaborado_por} — {ahora():%d/%m/%Y %H:%M}"


def a_csv(reporte: Reporte) -> bytes:
    buffer = io.StringIO()
    escritor = csv.writer(buffer)
    escritor.writerow([etiqueta for _, etiqueta in reporte.columnas])
    for fila in reporte.filas:
        escritor.writerow([fila.get(clave, "") if not isinstance(fila.get(clave), Decimal) else str(fila[clave])
                           for clave, _ in reporte.columnas])
    for totales in _filas_totales(reporte):
        escritor.writerow(["" if v is None else str(v) for v in totales])
    if pie := _pie(reporte):
        escritor.writerow([])
        escritor.writerow([pie])
    return buffer.getvalue().encode("utf-8-sig")  # BOM para que Excel respete acentos


def a_xlsx(reporte: Reporte) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = reporte.titulo[:31]
    hoja.append([etiqueta for _, etiqueta in reporte.columnas])
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for fila in reporte.filas:
        hoja.append([_celda_excel(fila.get(clave)) for clave, _ in reporte.columnas])
        _rojo_si_negativo(hoja[hoja.max_row])
    for totales in _filas_totales(reporte):
        hoja.append([_celda_excel(v) for v in totales])
        for celda in hoja[hoja.max_row]:
            celda.font = Font(bold=True)
        _rojo_si_negativo(hoja[hoja.max_row], negritas=True)
    for columna in hoja.columns:
        ancho = max(len(_texto(c.value)) for c in columna)
        hoja.column_dimensions[columna[0].column_letter].width = min(max(ancho + 2, 10), 60)
    if pie := _pie(reporte):  # después de calcular anchos, para no ensanchar la primera columna
        hoja.append([])
        hoja.append([pie])
        hoja.cell(row=hoja.max_row, column=1).font = Font(italic=True)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


ROJO = "C00000"


def _negativo(valor: Any) -> bool:
    return isinstance(valor, (Decimal, int, float)) and not isinstance(valor, bool) and valor < 0


def _rojo_si_negativo(celdas, negritas: bool = False) -> None:
    """Los saldos negativos (p. ej. un proyecto que usó dinero de otros) van en rojo."""
    for celda in celdas:
        if _negativo(celda.value):
            celda.font = Font(color=ROJO, bold=negritas)


def _celda_excel(valor: Any) -> Any:
    if isinstance(valor, datetime):
        return valor.replace(tzinfo=None)
    if isinstance(valor, (dict, list)):
        return str(valor)
    return valor


def _pdf(valor: Any) -> str:
    texto = escape(_texto(valor))
    return f'<font color="#{ROJO}">{texto}</font>' if _negativo(valor) else texto


def a_pdf(reporte: Reporte) -> bytes:
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(buffer, pagesize=landscape(letter), leftMargin=30, rightMargin=30)
    estilos = getSampleStyleSheet()
    celda = estilos["BodyText"].clone("celda", fontSize=7, leading=9)
    contenido: list[Any] = [Paragraph(reporte.titulo, estilos["Title"])]
    if reporte.parametros:
        filtros = ", ".join(f"{k}: {_texto(v)}" for k, v in reporte.parametros.items() if v is not None)
        contenido.append(Paragraph(filtros, estilos["Normal"]))
    contenido.append(Spacer(1, 10))

    datos = [[Paragraph(f"<b>{etiqueta}</b>", celda) for _, etiqueta in reporte.columnas]]
    for fila in reporte.filas:
        datos.append([Paragraph(_pdf(fila.get(clave)), celda) for clave, _ in reporte.columnas])
    for totales in _filas_totales(reporte):
        datos.append([Paragraph(f"<b>{_pdf(v)}</b>", celda) for v in totales])
    if len(datos) == 1:
        datos.append([Paragraph("Sin registros", celda)] + [""] * (len(reporte.columnas) - 1))

    tabla = Table(datos, repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.black),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    contenido.append(tabla)
    if pie := _pie(reporte):
        contenido += [Spacer(1, 14), Paragraph(f"<i>{pie}</i>", estilos["Normal"])]
    documento.build(contenido)
    return buffer.getvalue()


# ------------------------------------------------------------------ estado de ingresos y egresos
# Colores del formato "Resumen de ingresos y egresos".
AZUL = "0E1A33"
VERDE = "008000"
BEIGE = "F2F0EB"
GRIS = "D9D9D9"
FORMATO_NUMERO = '#,##0.00;-#,##0.00;"-"'


def _concepto(fila: dict[str, Any], separador: str = " · ") -> str:
    return f"{fila['concepto']}{separador}{fila['detalle']}" if fila.get("detalle") else fila["concepto"]


def _monedas(reporte: Reporte) -> list[str]:
    return [clave for clave, _ in reporte.columnas if clave != "concepto"]


def _xlsx_resumen(reporte: Reporte) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Resumen"
    monedas = _monedas(reporte)
    hoja.column_dimensions["A"].width = 75
    for n in range(2, 2 + len(monedas)):
        hoja.column_dimensions[get_column_letter(n)].width = 18

    hoja.append([reporte.titulo])
    hoja["A1"].font = Font(bold=True, size=16, color=AZUL)
    hoja.append([", ".join(f"{k}: {_texto(v)}" for k, v in reporte.parametros.items() if v is not None)])
    hoja["A2"].font = Font(italic=True, color="666666")
    hoja.append([])
    hoja.append(["Concepto", *monedas])
    for i, celda in enumerate(hoja[hoja.max_row]):
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        if i:
            celda.alignment = Alignment(horizontal="right")

    linea_fina = Border(bottom=Side(style="thin", color=GRIS))
    linea_total = Border(top=Side(style="medium", color=AZUL))
    for fila in reporte.filas:
        estilo = fila["estilo"]
        if estilo in ("seccion", "resultado") and hoja.max_row > 4:
            hoja.append([])  # espacio antes de cada sección y del resultado
        texto = _concepto(fila) if estilo in ("ingreso", "egreso") else fila["concepto"]
        hoja.append([texto, *[_celda_excel(fila.get(m)) for m in monedas]])
        celdas = hoja[hoja.max_row]
        for celda in celdas[1:]:
            celda.number_format = FORMATO_NUMERO
        if estilo == "seccion":
            celdas[0].font = Font(bold=True, size=12, color=AZUL)
        elif estilo in ("ingreso", "egreso"):
            for celda in celdas:
                celda.border = linea_fina
            for celda in celdas[1:]:
                if celda.value is not None:
                    color = ROJO if _negativo(celda.value) else (VERDE if estilo == "ingreso" else "000000")
                    celda.font = Font(color=color)
        elif estilo in ("total", "saldo"):
            for celda in celdas:
                celda.font = Font(bold=True, color=ROJO if _negativo(celda.value) else "000000")
                celda.fill = PatternFill("solid", fgColor=BEIGE)
                celda.border = linea_total
        elif estilo == "resultado":
            for celda in celdas:
                celda.font = Font(bold=True, color="FFC7CE" if _negativo(celda.value) else "FFFFFF")
                celda.fill = PatternFill("solid", fgColor=AZUL)
    if pie := _pie(reporte):
        hoja.append([])
        hoja.append([pie])
        hoja.cell(row=hoja.max_row, column=1).font = Font(italic=True)
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _pdf_resumen(reporte: Reporte) -> bytes:
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=36, rightMargin=36)
    estilos = getSampleStyleSheet()
    titulo = estilos["Title"].clone("titulo_resumen", alignment=0, textColor=colors.HexColor(f"#{AZUL}"))
    celda = estilos["BodyText"].clone("celda_resumen", fontSize=8, leading=10)
    derecha = celda.clone("derecha_resumen", alignment=2)
    monedas = _monedas(reporte)
    contenido: list[Any] = [Paragraph(escape(reporte.titulo), titulo)]
    if reporte.parametros:
        filtros = ", ".join(f"{k}: {_texto(v)}" for k, v in reporte.parametros.items() if v is not None)
        contenido.append(Paragraph(escape(filtros), estilos["Normal"]))
    contenido.append(Spacer(1, 10))

    def numero(valor: Any, color: str | None = None, negritas: bool = False) -> Paragraph:
        if valor is None:
            return Paragraph("", derecha)
        texto = "-" if valor == 0 else escape(_texto(valor))
        if _negativo(valor) and color != "FFFFFF":
            color = ROJO
        if color:
            texto = f'<font color="#{color}">{texto}</font>'
        return Paragraph(f"<b>{texto}</b>" if negritas else texto, derecha)

    blanco = celda.clone("blanco_resumen", textColor=colors.white)
    datos: list[list[Any]] = [[Paragraph("<b>Concepto</b>", blanco)] +
                              [Paragraph(f"<b>{m}</b>", blanco.clone(f"h{m}", alignment=2)) for m in monedas]]
    comandos: list[tuple] = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{AZUL}")),
                             ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    vacia = [""] * (1 + len(monedas))
    for fila in reporte.filas:
        estilo = fila["estilo"]
        if estilo in ("seccion", "resultado") and len(datos) > 1:
            datos.append(list(vacia))
        n = len(datos)
        if estilo == "seccion":
            datos.append([Paragraph(f'<font color="#{AZUL}"><b>{escape(fila["concepto"])}</b></font>', celda)]
                         + [""] * len(monedas))
        elif estilo in ("ingreso", "egreso"):
            texto = escape(fila["concepto"])
            if fila.get("detalle"):
                texto += f'<br/><font size="6.5" color="#666666">{escape(fila["detalle"])}</font>'
            datos.append([Paragraph(texto, celda)] +
                         [numero(fila.get(m), VERDE if estilo == "ingreso" else None) for m in monedas])
            comandos.append(("LINEBELOW", (0, n), (-1, n), 0.25, colors.HexColor(f"#{GRIS}")))
        elif estilo in ("total", "saldo"):
            datos.append([Paragraph(f"<b>{escape(fila['concepto'])}</b>", celda)] +
                         [numero(fila.get(m), negritas=True) for m in monedas])
            comandos += [("BACKGROUND", (0, n), (-1, n), colors.HexColor(f"#{BEIGE}")),
                         ("LINEABOVE", (0, n), (-1, n), 1, colors.HexColor(f"#{AZUL}"))]
        elif estilo == "resultado":
            datos.append([Paragraph(f"<b>{escape(fila['concepto'])}</b>", blanco)] +
                         [numero(fila.get(m), "FFC7CE" if _negativo(fila.get(m)) else "FFFFFF", True)
                          for m in monedas])
            comandos.append(("BACKGROUND", (0, n), (-1, n), colors.HexColor(f"#{AZUL}")))
    ancho_moneda = 95
    tabla = Table(datos, colWidths=[documento.width - ancho_moneda * len(monedas)] + [ancho_moneda] * len(monedas),
                  repeatRows=1)
    tabla.setStyle(TableStyle(comandos))
    contenido.append(tabla)
    if pie := _pie(reporte):
        contenido += [Spacer(1, 14), Paragraph(f"<i>{escape(pie)}</i>", estilos["Normal"])]
    documento.build(contenido)
    return buffer.getvalue()


def _csv_resumen(reporte: Reporte) -> bytes:
    buffer = io.StringIO()
    escritor = csv.writer(buffer)
    monedas = _monedas(reporte)
    escritor.writerow(["Concepto", *monedas])
    for fila in reporte.filas:
        escritor.writerow([_concepto(fila), *["" if fila.get(m) is None else str(fila[m]) for m in monedas]])
    if pie := _pie(reporte):
        escritor.writerow([])
        escritor.writerow([pie])
    return buffer.getvalue().encode("utf-8-sig")


def exportar(reporte: Reporte, formato: Formato) -> bytes:
    if reporte.estilo == "resumen":
        return {Formato.CSV: _csv_resumen, Formato.XLSX: _xlsx_resumen, Formato.PDF: _pdf_resumen}[formato](reporte)
    return {Formato.CSV: a_csv, Formato.XLSX: a_xlsx, Formato.PDF: a_pdf}[formato](reporte)
