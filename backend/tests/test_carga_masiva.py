import io
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from tests.helpers import crear_cuenta, crear_empresa, crear_proyecto, entrada, ok, saldo_cuenta, saldos_proyecto


def _excel(filas, encabezado=("FECHA", "CONCEPTO", "CARGO", "ABONO")) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.append(list(encabezado))
    for fila in filas:
        hoja.append(list(fila))
    buffer = io.BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


# Como el archivo del cliente: fechas de Excel, salidas antes que entradas y filas vacías al final.
CLIENTE = [
    (datetime(2026, 1, 7), "INGRESO A", None, 100000.0),
    (datetime(2026, 1, 6), "EGRESO A", 20000.0, None),
    (datetime(2026, 1, 8), "INGRESO B", None, "50,000.50"),
    (datetime(2026, 1, 9), "EGRESO B", 30000, None),
    (None, None, None, None),
]


@pytest.fixture
def base(client, juan):
    h = juan[1]
    empresa = crear_empresa(client, h)
    return {"h": h, "empresa": empresa, "banco": crear_cuenta(client, h, empresa["id"]),
            "caja": crear_cuenta(client, h, empresa["id"], "caja"), "proyecto": crear_proyecto(client, h)}


def _enviar(client, base, contenido, confirmar=False, salidas=None, **extra):
    datos = {"empresa_id": base["empresa"]["id"], "cuenta_entradas_id": base["banco"]["id"],
             "cuenta_salidas_id": (salidas or base["banco"])["id"], **extra}
    ruta = f"/proyectos/{base['proyecto']['id']}/carga-masiva" + ("" if confirmar else "/vista-previa")
    return client.post(ruta, data=datos, headers=base["h"],
                       files={"archivo": ("PLANTILLA.xlsx", contenido, "application/octet-stream")})


def test_vista_previa_no_registra_nada(client, base):
    previa = ok(_enviar(client, base, _excel(CLIENTE), porcentaje_gasto="10"))
    assert previa["valido"] and previa["errores"] == []
    assert (previa["entradas"], previa["salidas"]) == (2, 2)
    assert previa["entradas_bruto"] == "150000.50" and previa["entradas_gasto"] == "15000.05"
    assert previa["total_salidas"] == "50000.00"
    assert previa["saldos"] == [{"cuenta": base["banco"]["nombre"], "moneda": "MXN", "antes": "0.00",
                                 "despues": "85000.45"}]
    # Primero las entradas y después las salidas, con el número de fila del Excel.
    assert [(f["fila"], f["tipo"]) for f in previa["filas"]] == [(2, "entrada"), (4, "entrada"), (3, "salida"),
                                                                 (5, "salida")]
    assert saldo_cuenta(client, base["h"], base["banco"]) == "0.00"
    assert ok(client.get("/entradas", headers=base["h"]))["total"] == 0


def test_carga_entradas_y_despues_salidas(client, base, admin):
    h = base["h"]
    categoria = ok(client.post("/categorias", json={"nombre": "Materiales"}, headers=h), 201)
    # Las salidas (50,000) solo alcanzan si antes se registran las entradas (150,000.50).
    resultado = ok(_enviar(client, base, _excel(CLIENTE), confirmar=True, categoria_salidas_id=categoria["id"],
                           requiere_factura_salidas="true"))
    assert resultado == {"entradas": 2, "salidas": 2}
    assert saldo_cuenta(client, h, base["banco"]) == "100000.50"
    assert saldos_proyecto(client, h, base["proyecto"])["MXN"]["saldo"] == "100000.50"
    salidas = ok(client.get("/salidas", headers=h))["items"]
    assert {s["concepto"] for s in salidas} == {"EGRESO A", "EGRESO B"}
    assert all(s["categoria_nombre"] == "Materiales" and s["requiere_factura"] and s["forma"] == "transferencia"
               for s in salidas)
    proyecto = ok(client.get(f"/proyectos/{base['proyecto']['id']}", headers=h))
    assert [e["id"] for e in proyecto["empresas"]] == [base["empresa"]["id"]]  # se vinculó sola
    bitacora = ok(client.get("/bitacora", params={"entidad": "Salida"}, headers=admin[1]))["items"]
    assert bitacora[0]["detalle"]["origen"] == "carga masiva"


def test_estructura_incorrecta_no_registra_nada(client, base):
    previa = ok(_enviar(client, base, _excel(CLIENTE, encabezado=("FECHA", "DESCRIPCION", "DEBE", "HABER"))))
    assert not previa["valido"] and "estructura esperada" in previa["errores"][0]["mensaje"]
    r = _enviar(client, base, b"no es excel", confirmar=True)
    assert r.status_code == 422, r.text
    assert ok(client.get("/entradas", headers=base["h"]))["total"] == 0


def test_una_fila_con_error_no_registra_nada(client, base):
    filas = CLIENTE[:4] + [
        (datetime(2026, 1, 10), "AMBOS", 10, 10),
        ("31/02/2026", "FECHA MALA", None, 5),
        (datetime(2026, 1, 11), None, 7, None),
        (datetime(2026, 1, 12), "TEXTO", "abc", None),
        (datetime(2099, 1, 1), "FUTURA", None, 1),
        (datetime(2026, 1, 13), "SIN MONTO", None, None),
    ]
    previa = ok(_enviar(client, base, _excel(filas)))
    assert not previa["valido"]
    errores = {(e["fila"], e["columna"]) for e in previa["errores"]}
    assert errores == {(6, "CARGO / ABONO"), (7, "FECHA"), (8, "CONCEPTO"), (9, "CARGO"), (10, "FECHA"),
                       (11, "CARGO / ABONO")}
    assert _enviar(client, base, _excel(filas), confirmar=True).status_code == 422
    assert ok(client.get("/entradas", headers=base["h"]))["total"] == 0


def test_cuenta_que_quedaria_en_negativo(client, base):
    # Entradas al banco, salidas de la caja vacía: la caja no alcanza y no se registra nada.
    previa = ok(_enviar(client, base, _excel(CLIENTE), salidas=base["caja"]))
    assert not previa["valido"] and "no alcanza" in previa["errores"][0]["mensaje"]
    assert _enviar(client, base, _excel(CLIENTE), confirmar=True, salidas=base["caja"]).status_code == 422
    assert saldo_cuenta(client, base["h"], base["banco"]) == "0.00"
    # Con dinero previo en la caja sí alcanza; las salidas de caja son en efectivo.
    entrada(client, base["h"], base["proyecto"], base["caja"], "60000", fecha="2026-01-01")
    ok(_enviar(client, base, _excel(CLIENTE), confirmar=True, salidas=base["caja"]))
    assert saldo_cuenta(client, base["h"], base["caja"]) == "10000.00"
    assert {s["forma"] for s in ok(client.get("/salidas", headers=base["h"]))["items"]} == {"efectivo"}


def test_cuenta_de_otra_empresa(client, base):
    otra = crear_empresa(client, base["h"], "Otra")
    ajena = crear_cuenta(client, base["h"], otra["id"])
    r = _enviar(client, base, _excel(CLIENTE), salidas=ajena)
    assert r.status_code == 422 and "no es de la empresa" in r.text


def test_plantilla(client, base):
    r = client.get("/carga-masiva/plantilla", headers=base["h"])
    assert r.status_code == 200
    hoja = load_workbook(io.BytesIO(r.content)).worksheets[0]
    assert [c.value for c in hoja[1]] == ["FECHA", "CONCEPTO", "CARGO", "ABONO"]
